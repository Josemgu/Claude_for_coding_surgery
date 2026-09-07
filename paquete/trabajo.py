"""La hoja «Por verificar» que se le entrega a un companero, construida en memoria.

Este modulo no toca el disco: construye el `Workbook` y lo devuelve. Escribirlo es
cosa de `paquete.exportacion`, igual que en `espejo/`.

**Es la hoja del proyecto viejo** (`salida/asignacion.py`), que es la que el dueno
llama perfecta (`DECISIONES.md`, 2026-09-03). Lo que se copia de alli, y por que
cada cosa:

  - **Cinco filas de cabecera antes de la tabla.** Lo que es igual para todas las
    filas se dice una vez arriba y no se repite en cada renglon: el caso, el
    templo, la fecha de salida, la fecha limite y a quien se le manda.
  - **La fecha limite sola y en rojo.** Es lo unico de la hoja que no se negocia:
    pasada esa fecha, quien no este listo viaja y no hace la ordenanza. Va una
    semana antes del viaje porque si falta algo hace falta tiempo para hablar con
    el lider.
  - **La clave A LA VISTA, en gris pequeno, en la ultima columna.** Citado del
    viejo: «un dato oculto es un dato que alguien borra sin saber lo que hace».
  - **Fondo `PIEL` en lo que rellena el companero** y menu desplegable de dos
    opciones. El menu es una ayuda, no una reja: una respuesta escrita a mano que
    no este en la lista se lee igual al volver, y si no se entiende se avisa con su
    numero de fila.

Sin `pandas` (regla permanente 3): solo `openpyxl`. Y en espanol todo lo que se ve
(regla permanente 4), porque lo lee una persona que no es Miguel.

**Tres cosas medidas en esta maquina con openpyxl 3.1.5:**

  1. **`numero_caso`, `mrn` y `clave` quedan bloqueadas contra edicion.** En OOXML
     el bloqueo de una celda no hace nada por si solo: solo surte efecto cuando la
     HOJA esta protegida. Por eso se hacen las dos cosas —`hoja.protection.sheet =
     True` y `Protection(locked=...)` celda a celda— y por eso todas las demas
     celdas se desbloquean explicitamente: por defecto **todas** las celdas nacen
     bloqueadas, asi que proteger la hoja sin desbloquear el resto dejaria al
     companero sin poder escribir nada.

     ⚠️ **Que NO es esto.** Es una barrera contra el error, no contra alguien que
     quiera saltarsela: la proteccion de hoja de Excel se quita desde el menu. Lo
     que impide es el accidente —teclear encima del MRN creyendo que se corrige—,
     que es el que de verdad pasa. La defensa de verdad esta en la vuelta: la fila
     cuya clave no casa no se inserta.

  2. ⚠️ `showDropDown` significa lo CONTRARIO de lo que parece. En openpyxl es un
     alias de `hide_drop_down` —comprobado en el codigo de la clase
     `DataValidation` de openpyxl 3.1.5—, asi que se deja en su valor por defecto
     `False` para que la flechita SE VEA. Poner `True` la esconderia.

  3. **El MRN y la clave salen con formato de texto.** `FORMAT_TEXT` es `'@'`. Sin
     el, Excel se come los ceros de delante de `055-1111-3853` y ese MRN ya no casa
     al volver.
"""

from collections import namedtuple
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from datos.pasos import RESPUESTAS
from paquete.columnas import (
    COLUMNAS,
    COLUMNA_DE_LA_CLAVE,
    FILA_DE_LA_CABECERA,
    NOMBRE_DE_LA_HOJA,
    PRIMERA_FILA_DE_DATOS,
    SIN_DATO,
    TEMPORAL,
    TEXTO,
    ancho_de,
    titulos,
)

_FORMATO_DE_FECHA = "%Y-%m-%d"
_FORMATO_QUE_SE_LEE = "%d-%m-%Y"

# La recomendacion tiene que estar lista antes de que salga el grupo, no el mismo
# dia: si falta algo, hace falta tiempo para hablar con el lider. Son los mismos
# siete dias del proyecto viejo.
DIAS_DE_MARGEN = 7

TINTA = "16233A"
PIEL = "FFF6DC"       # lo que rellena el companero
LINEA = "DBDAD2"
ROJO = "A62E24"       # solo la fecha limite
GRIS_DEL_SUBTITULO = "4B5872"
GRIS_DE_LA_CLAVE = "8A93A5"

TAMANO_DEL_TITULO = 14
TAMANO_DE_LA_CLAVE = 8
ALTO_DE_LA_INSTRUCCION = 30
ALTO_DE_LA_CABECERA = 32

TITULO = "Preparación para las ordenanzas"
INSTRUCCION = (
    "Busca a cada persona en el sistema, entra en «Preparación para las ordenanzas» "
    "y marca Sí o No en cada paso, tal como lo veas. Si algo no está completo, llama "
    "al líder de su barrio y ayúdalo a terminarlo."
)
SIN_AGENTE = "sin asignar"

# Lo que va arriba de la tabla y es igual para todas sus filas.
Cabecera = namedtuple("Cabecera", ("numero_de_caso", "templo", "fecha_de_salida", "agente"))


class ErrorDelPaquete(RuntimeError):
    """El paquete no se pudo construir con lo que se le dio."""


def _fecha_legible(iso):
    """Una fecha ISO escrita como la lee una persona, o vacio si no se puede."""
    try:
        return datetime.strptime(iso, _FORMATO_DE_FECHA).strftime(_FORMATO_QUE_SE_LEE)
    except (ValueError, TypeError):
        return ""


def fecha_limite(fecha_de_salida):
    """Una semana antes del viaje, escrita como la lee una persona.

    Devuelve vacio si la fecha de salida no se puede leer. No se inventa un margen
    sobre una fecha que no se entiende: una fecha limite falsa es peor que ninguna,
    porque el companero se organiza contra ella.
    """
    try:
        salida = datetime.strptime(fecha_de_salida, _FORMATO_DE_FECHA).date()
    except (ValueError, TypeError):
        return ""
    return (salida - timedelta(days=DIAS_DE_MARGEN)).strftime(_FORMATO_QUE_SE_LEE)


def cabecera_de(filas, agente=""):
    """Lo que va arriba de la tabla, sacado de las propias filas.

    El caso se nombra solo cuando todas las filas son del mismo, igual que en el
    viejo: un titulo que dice un caso concreto sobre una hoja que lleva tres es
    peor que un titulo generico.

    La fecha de salida es la **mas temprana** de las que traiga la hoja, y de ella
    sale la fecha limite. Con varios casos en el mismo paquete, la que manda es la
    del primero que viaja: una fecha limite calculada sobre el ultimo dejaria pasar
    sin aviso al grupo que sale antes.
    """
    casos = {fila.get("numero_caso") for fila in filas if fila.get("numero_caso")}
    salidas = sorted(
        fila["fecha_viaje"] for fila in filas if fila.get("fecha_viaje")
    )
    templos = {fila.get("templo") for fila in filas if fila.get("templo")}
    return Cabecera(
        numero_de_caso=next(iter(casos)) if len(casos) == 1 else "",
        templo=next(iter(templos)) if len(templos) == 1 else "",
        fecha_de_salida=salidas[0] if salidas else "",
        agente=agente or "",
    )


def _escribir_una_linea(hoja, numero_de_fila, texto, fuente=None):
    """Una de las cinco lineas de arriba, siempre en la columna A."""
    celda = hoja.cell(row=numero_de_fila, column=1, value=texto)
    if fuente is not None:
        celda.font = fuente
    return celda


def _escribir_las_cinco_lineas(hoja, cabecera):
    """Titulo, templo y salida, fecha limite, agente e instruccion. En ese orden.

    La linea del templo se arma juntando solo los trozos que existen, igual que en
    el viejo: sin templo guardado, sale «Sale el ...» a secas en vez de «Templo:  ·
    Sale el ...», que parece un dato que se perdio.

    La fecha limite se escribe **solo si se puede calcular**, y por eso la fila 3
    puede quedar vacia. Es preferible a escribir «Todo verificado antes del » sin
    fecha detras, que no dice nada y ocupa el sitio de lo que si importa.
    """
    titulo = TITULO + (f" · {cabecera.numero_de_caso}" if cabecera.numero_de_caso else "")
    _escribir_una_linea(
        hoja, 1, titulo, Font(bold=True, size=TAMANO_DEL_TITULO, color=TINTA)
    )

    trozos = [
        f"Templo: {cabecera.templo}" if cabecera.templo else "",
        f"Sale el {_fecha_legible(cabecera.fecha_de_salida)}"
        if _fecha_legible(cabecera.fecha_de_salida)
        else "",
    ]
    _escribir_una_linea(
        hoja, 2, " · ".join(trozo for trozo in trozos if trozo),
        Font(color=GRIS_DEL_SUBTITULO),
    )

    limite = fecha_limite(cabecera.fecha_de_salida)
    if limite:
        _escribir_una_linea(
            hoja, 3, f"Todo verificado antes del {limite}", Font(bold=True, color=ROJO)
        )

    _escribir_una_linea(
        hoja, 4, f"Agente: {cabecera.agente or SIN_AGENTE}", Font(bold=True, color=TINTA)
    )

    celda = _escribir_una_linea(hoja, 5, INSTRUCCION)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.row_dimensions[5].height = ALTO_DE_LA_INSTRUCCION


def _escribir_los_titulos(hoja):
    """La fila 6: fondo tinta, letra blanca, bloqueada entera y congelada debajo."""
    for numero, titulo in enumerate(titulos(), start=1):
        celda = hoja.cell(row=FILA_DE_LA_CABECERA, column=numero, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=TINTA)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        celda.protection = Protection(locked=True)
    hoja.row_dimensions[FILA_DE_LA_CABECERA].height = ALTO_DE_LA_CABECERA


def _valor_de_la_celda(valor, clase):
    """El valor que entra en la celda, segun la clase de su columna.

    Una fecha que no se puede leer se deja tal cual. No se adivina ni se descarta:
    la regla permanente 1 prohibe inventar un dato, y una fecha inventada en el
    papel que lleva un companero es peor que una fecha fea.
    """
    if valor is None:
        return None
    if clase == TEXTO:
        return str(valor)
    if clase == TEMPORAL and isinstance(valor, str):
        try:
            return datetime.strptime(valor, _FORMATO_DE_FECHA).date()
        except ValueError:
            return valor
    return valor


def _escribir_celda(hoja, fila, numero_de_columna, columna, valor):
    """Escribe una celda con su valor, su formato, su fondo y su bloqueo."""
    celda = hoja.cell(
        row=fila, column=numero_de_columna, value=_valor_de_la_celda(valor, columna.clase)
    )
    if columna.clase == TEXTO:
        celda.number_format = FORMAT_TEXT
    if celda.data_type == "f":
        # Un texto que empieza por `=` lo escribiria openpyxl como FORMULA. Un
        # nombre leido por OCR no es una formula. Mismo cuidado que en
        # `espejo/libro.py`.
        celda.data_type = "s"
    celda.protection = Protection(locked=not columna.editable)
    celda.border = Border(bottom=Side(style="thin", color=LINEA))
    if columna.respuesta:
        celda.fill = PatternFill("solid", fgColor=PIEL)
    if columna.nombre == COLUMNA_DE_LA_CLAVE:
        celda.font = Font(size=TAMANO_DE_LA_CLAVE, color=GRIS_DE_LA_CLAVE)
    return celda


def _poner_los_menus(hoja, ultima_fila):
    """Un menu de dos opciones en cada columna que rellena el companero.

    Uno por columna y no uno compartido: openpyxl escribe cada `DataValidation` con
    los rangos que se le anaden, y repartirlos en siete deja el archivo legible si
    alguien lo abre por dentro para ver por que una columna dejo de ofrecer el menu.

    `showErrorMessage` se queda en `False` **a proposito**, al reves que en el Excel
    de trabajo anterior. El menu es una ayuda, no una reja: el proyecto viejo lo
    dice con esas palabras y lo sostiene la vuelta, que lee una respuesta escrita a
    mano igual y avisa con su numero de fila si no la entiende. Bloquear la celda
    obligaria al companero a dejarla vacia cuando la realidad no cabe en dos
    opciones, y una celda vacia es peor: no distingue «no aplica» de «no lo mire».
    """
    if ultima_fila < PRIMERA_FILA_DE_DATOS:
        return ()
    puestos = []
    for numero, columna in enumerate(COLUMNAS, start=1):
        if not columna.respuesta:
            continue
        validacion = DataValidation(
            type="list",
            formula1='"' + ",".join(RESPUESTAS) + '"',
            allow_blank=True,
        )
        hoja.add_data_validation(validacion)
        letra = get_column_letter(numero)
        validacion.add(f"{letra}{PRIMERA_FILA_DE_DATOS}:{letra}{ultima_fila}")
        puestos.append(validacion)
    return tuple(puestos)


def _ensanchar_las_columnas(hoja):
    """Deja cada columna con el ancho de su contenido. Un MRN estrecho sale `#####`."""
    for numero, columna in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(numero)].width = ancho_de(columna.nombre)


def construir_libro_de_trabajo(filas, agente="", cabecera=None):
    """El libro que se le entrega al companero. `filas` son diccionarios de persona.

    Cada fila necesita, como minimo, las claves que nombra `paquete.columnas`. Lo
    que falte se escribe con la palabra que dice que no consta: es preferible a no
    entregar el paquete, y se ve.

    `cabecera` se deduce de las propias filas si no se pasa, para que quien solo
    tenga las filas —una prueba, una llamada suelta— no tenga que armarla aparte.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = NOMBRE_DE_LA_HOJA

    filas = list(filas)
    _escribir_las_cinco_lineas(hoja, cabecera or cabecera_de(filas, agente))
    _escribir_los_titulos(hoja)

    numero_de_fila = FILA_DE_LA_CABECERA
    for numero_de_fila, fila in enumerate(filas, start=PRIMERA_FILA_DE_DATOS):
        for numero_de_columna, columna in enumerate(COLUMNAS, start=1):
            valor = fila.get(columna.nombre)
            if valor is None and not columna.respuesta:
                # Lo que el programa no sabe se dice con palabras y no con un
                # hueco. Un hueco lo lee el companero como «esto lo relleno yo», y
                # las unicas celdas que rellena el son las de fondo amarillo.
                valor = SIN_DATO
            _escribir_celda(hoja, numero_de_fila, numero_de_columna, columna, valor)

    _poner_los_menus(hoja, numero_de_fila)
    _ensanchar_las_columnas(hoja)
    hoja.freeze_panes = f"A{PRIMERA_FILA_DE_DATOS}"
    # La proteccion de la hoja va DESPUES de escribirlo todo. Es lo que activa los
    # bloqueos de celda de arriba; sin esta linea, `locked=True` no impide nada.
    hoja.protection.sheet = True
    return libro
