"""La hoja «Por verificar»: que salga y vuelva como la del proyecto viejo.

Las pruebas de aquí salen de los criterios de cierre 1 y 2 del pase del
2026-09-03, no de leer el código:

  1. «El Excel del agente del nuevo tiene la misma hoja, las mismas cabeceras de
     las filas 1 a 5, las mismas columnas en el mismo orden, los menús y la clave
     visible que el del viejo.»
  2. «Cargar ese Excel de vuelta con respuestas —incluida una fuera del menú y una
     fila en blanco— cae en la persona correcta y avisa como el viejo.»

⚠️ **De dónde salen los literales de este archivo.** La estructura del viejo se
midió abriendo con `openpyxl` una hoja generada por
`C:\\Users\\josem\\Desktop\\proyecto\\pdf-a-excel\\salida\\asignacion.py` con datos
inventados, el 2026-09-03. Los valores medidos se escriben aquí **como literales**
y no se leen de aquel proyecto: una prueba que importa código de una carpeta que no
está en este repositorio deja de pasar en cuanto alguien mueve esa carpeta, y
entonces lo que falla no es el programa.

Lo medido, literal, en aquella hoja:

    hoja                      «Por verificar», y es la única
    fila de títulos           6, congelada en A7
    columnas                  Caso · Fecha de solicitud · Fecha de viaje ·
                              Barrio o rama · Estaca o distrito ·
                              Hermano(a) que viaja · Cédula de miembro · A qué va ·
                              1. Preparación … 6. Listo para el templo ·
                              ¿Llamó al líder? · clave
    fondo de la cabecera      16233A, letra FFFFFF
    fondo de lo que rellena   FFF6DC
    la clave                  visible, última columna, tamaño 8, color 8A93A5
    menús                     7, de tipo lista, con dos opciones
"""

import unittest

from openpyxl import load_workbook

from datos.asignaciones import asignar_casos
from datos.companeros import alta_de_companero
from datos.pasos import COLUMNA_DE_LA_LLAMADA, NOMBRES_DE_LOS_PASOS, RESPUESTAS
from datos.propuestas import resolver_persona_por_par
from datos.repositorio import alta_de_caso, alta_de_persona
from paquete.columnas import (
    FILA_DE_LA_CABECERA,
    NOMBRE_DE_LA_HOJA,
    PRIMERA_FILA_DE_DATOS,
    SIN_DATO,
    indice_de,
    titulos,
)
from paquete.exportacion import exportar_paquete
from paquete.reconciliacion import reconciliar_excel
from pruebas.comun import PruebaConBaseTemporal

# ---- lo medido en la hoja del proyecto viejo el 2026-09-03 -------------------

HOJA_DEL_VIEJO = "Por verificar"
FILA_DE_TITULOS_DEL_VIEJO = 6
CONGELADO_EN_EL_VIEJO = "A7"

TITULOS_DEL_VIEJO = [
    "Caso",
    "Fecha de solicitud",
    "Fecha de viaje",
    "Barrio o rama",
    "Estaca o distrito",
    "Hermano(a) que viaja",
    "Cédula de miembro",
    "A qué va",
    "1. Preparación",
    "2. Información",
    "3. Cita del templo",
    "4. Acciones requeridas",
    "5. Entrevistas",
    "6. Listo para el templo",
    "¿Llamó al líder?",
    "clave",
]

TINTA_DEL_VIEJO = "0016233A"
BLANCO = "00FFFFFF"
PIEL_DEL_VIEJO = "00FFF6DC"
GRIS_DE_LA_CLAVE_DEL_VIEJO = "008A93A5"
TAMANO_DE_LA_CLAVE_DEL_VIEJO = 8
ROJO_DE_LA_FECHA_LIMITE_DEL_VIEJO = "00A62E24"

FECHA_DE_VIAJE = "2026-10-20"
# Siete días antes, que es el margen del viejo, escrito como lo lee una persona.
FECHA_LIMITE_ESPERADA = "13-10-2026"


class BaseConUnPaquete(PruebaConBaseTemporal):
    """Un caso con dos personas, asignado a un compañero, y su paquete escrito."""

    def setUp(self):
        super().setUp()
        self.companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        self.caso_id = alta_de_caso(
            self.conexion,
            "CASP2610",
            unidad_numero="123456",
            unidad_nombre="Paramaribo Branch",
            fecha_viaje=FECHA_DE_VIAJE,
        ).id
        self.personas = [
            alta_de_persona(
                self.conexion,
                self.caso_id,
                mrn=f"055-1111-385{fila}",
                nombre=f"Persona {fila}",
                fila_formulario=fila,
                ord_investidura=1 if fila == 1 else None,
            )
            for fila in (1, 2)
        ]
        asignar_casos(self.conexion, [self.caso_id], self.companero_id)
        self.resultado = exportar_paquete(
            self.conexion, self.companero_id, self.carpeta_temporal
        )

    def abrir(self):
        """La hoja escrita, tal como la abriría el compañero."""
        libro = load_workbook(str(self.resultado.ruta_del_excel))
        return libro, libro[NOMBRE_DE_LA_HOJA]

    def editar(self, cambios):
        """Edita la hoja POR FUERA, como haría el compañero, y la guarda."""
        libro = load_workbook(str(self.resultado.ruta_del_excel))
        cambios(libro[NOMBRE_DE_LA_HOJA])
        libro.save(str(self.resultado.ruta_del_excel))
        libro.close()

    def fila_de(self, hoja, mrn):
        for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
            if hoja.cell(row=fila, column=indice_de("mrn")).value == mrn:
                return fila
        raise AssertionError(f"No hay ninguna fila con el MRN {mrn}.")


class PruebaDeLaFormaDeLaHoja(BaseConUnPaquete):
    """Criterio 1: la misma hoja, las mismas cabeceras y las mismas columnas."""

    def test_hay_una_sola_hoja_y_se_llama_como_la_del_viejo(self):
        libro, hoja = self.abrir()
        nombres = libro.sheetnames
        libro.close()
        self.assertEqual([HOJA_DEL_VIEJO], nombres)
        self.assertEqual(HOJA_DEL_VIEJO, hoja.title)

    def test_las_columnas_son_las_mismas_y_en_el_mismo_orden(self):
        """El orden importa tanto como los nombres: el agente lee de izquierda a derecha."""
        self.assertEqual(TITULOS_DEL_VIEJO, titulos())

        libro, hoja = self.abrir()
        escritos = [
            hoja.cell(row=FILA_DE_TITULOS_DEL_VIEJO, column=numero).value
            for numero in range(1, len(TITULOS_DEL_VIEJO) + 1)
        ]
        libro.close()
        self.assertEqual(TITULOS_DEL_VIEJO, escritos)

    def test_la_tabla_empieza_en_la_fila_6_y_se_congela_ahi(self):
        libro, hoja = self.abrir()
        congelado = hoja.freeze_panes
        libro.close()
        self.assertEqual(FILA_DE_TITULOS_DEL_VIEJO, FILA_DE_LA_CABECERA)
        self.assertEqual(CONGELADO_EN_EL_VIEJO, congelado)

    def test_las_cinco_lineas_de_arriba_dicen_lo_que_decian_en_el_viejo(self):
        """Título con el caso, la salida, la fecha límite en rojo, el agente y la instrucción."""
        libro, hoja = self.abrir()
        textos = {
            numero: hoja.cell(row=numero, column=1).value for numero in range(1, 6)
        }
        libro.close()

        self.assertIn("Preparación para las ordenanzas", textos[1])
        self.assertIn("CASP2610", textos[1])
        self.assertIn("Sale el 20-10-2026", textos[2])
        self.assertEqual(f"Todo verificado antes del {FECHA_LIMITE_ESPERADA}", textos[3])
        self.assertEqual("Agente: Ana Pérez", textos[4])
        self.assertIn("marca Sí o No en cada paso", textos[5])

    def test_la_fecha_limite_va_sola_y_en_rojo(self):
        """Es lo único de la hoja que no se negocia, y por eso se ve distinto."""
        libro, hoja = self.abrir()
        celda = hoja.cell(row=3, column=1)
        color = celda.font.color.rgb
        negrita = celda.font.bold
        libro.close()
        self.assertEqual(ROJO_DE_LA_FECHA_LIMITE_DEL_VIEJO, color)
        self.assertTrue(negrita)

    def test_la_cabecera_va_en_tinta_con_letra_blanca(self):
        libro, hoja = self.abrir()
        celda = hoja.cell(row=FILA_DE_LA_CABECERA, column=1)
        fondo = celda.fill.fgColor.rgb
        letra = celda.font.color.rgb
        libro.close()
        self.assertEqual(TINTA_DEL_VIEJO, fondo)
        self.assertEqual(BLANCO, letra)

    def test_solo_lo_que_rellena_el_agente_lleva_fondo_piel(self):
        """Siete columnas en amarillo y ni una más: el fondo es la instrucción.

        El nombre queda editable pero SIN fondo. Pintarlo le pediría al compañero
        que teclease un nombre que ya viene puesto.
        """
        libro, hoja = self.abrir()
        con_piel = {
            hoja.cell(row=FILA_DE_LA_CABECERA, column=numero).value
            for numero in range(1, len(TITULOS_DEL_VIEJO) + 1)
            if hoja.cell(row=PRIMERA_FILA_DE_DATOS, column=numero).fill.fgColor.rgb
            == PIEL_DEL_VIEJO
        }
        libro.close()

        esperadas = set(TITULOS_DEL_VIEJO[8:15])
        self.assertEqual(esperadas, con_piel)
        self.assertNotIn("Hermano(a) que viaja", con_piel)

    def test_la_clave_es_visible_en_la_ultima_columna_y_en_gris_pequeno(self):
        """«Un dato oculto es un dato que alguien borra sin saber lo que hace»."""
        libro, hoja = self.abrir()
        columna = indice_de("clave")
        celda = hoja.cell(row=PRIMERA_FILA_DE_DATOS, column=columna)
        valor, tamano, color = celda.value, celda.font.size, celda.font.color.rgb
        oculta = hoja.column_dimensions[celda.column_letter].hidden
        libro.close()

        self.assertEqual(len(TITULOS_DEL_VIEJO), columna, "la clave va la última")
        self.assertFalse(oculta, "la clave NO se esconde")
        # Tres partes desde el 2026-09-03: número de caso, MRN e **id del caso**.
        # Sin el id, dos casos que comparten número —y desde la versión 12 pueden—
        # devolverían la misma clave y la fila que vuelve no sabría a cuál es.
        self.assertEqual(f"CASP2610:055-1111-3851:{self.caso_id}", valor)
        self.assertEqual(TAMANO_DE_LA_CLAVE_DEL_VIEJO, tamano)
        self.assertEqual(GRIS_DE_LA_CLAVE_DEL_VIEJO, color)

    def test_las_siete_columnas_del_agente_traen_su_menu_de_si_y_no(self):
        libro, hoja = self.abrir()
        formulas = {v.formula1 for v in hoja.data_validations.dataValidation}
        cuantos = len(hoja.data_validations.dataValidation)
        libro.close()

        self.assertEqual(len(NOMBRES_DE_LOS_PASOS) + 1, cuantos)
        self.assertEqual({'"' + ",".join(RESPUESTAS) + '"'}, formulas)

    def test_lo_que_la_base_no_guarda_se_dice_con_palabras_y_no_en_blanco(self):
        """«Estaca o distrito» y «Fecha de solicitud» no están en la base.

        Salen con la palabra que dice que no consta y no vacías: un hueco lo lee el
        compañero como «esto lo relleno yo», y lo que rellena él es lo amarillo.
        """
        libro, hoja = self.abrir()
        valores = {
            nombre: hoja.cell(row=PRIMERA_FILA_DE_DATOS, column=indice_de(nombre)).value
            for nombre in ("estaca", "fecha_solicitud")
        }
        libro.close()
        self.assertEqual({"estaca": SIN_DATO, "fecha_solicitud": SIN_DATO}, valores)

    def test_a_que_va_sale_escrito_en_espanol_y_no_como_nombre_de_columna(self):
        libro, hoja = self.abrir()
        con_marca = hoja.cell(
            row=self.fila_de(hoja, "055-1111-3851"), column=indice_de("a_que_va")
        ).value
        sin_marca = hoja.cell(
            row=self.fila_de(hoja, "055-1111-3852"), column=indice_de("a_que_va")
        ).value
        libro.close()
        self.assertEqual("Investidura", con_marca)
        self.assertEqual("sin marcar", sin_marca)


class PruebaDeLaVueltaDeLaHoja(BaseConUnPaquete):
    """Criterio 2: cae en la persona correcta y avisa como el viejo."""

    def test_las_respuestas_caen_en_la_persona_de_su_fila(self):
        def contestar(hoja):
            primera = self.fila_de(hoja, "055-1111-3851")
            segunda = self.fila_de(hoja, "055-1111-3852")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=primera, column=indice_de(nombre)).value = "Sí"
                hoja.cell(row=segunda, column=indice_de(nombre)).value = "No"
            hoja.cell(row=primera, column=indice_de(COLUMNA_DE_LA_LLAMADA)).value = "No"
            hoja.cell(row=segunda, column=indice_de(COLUMNA_DE_LA_LLAMADA)).value = "Sí"

        self.editar(contestar)
        vuelta = reconciliar_excel(
            self.conexion, self.resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(2, len(vuelta.aplicadas))
        primera = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3851")
        segunda = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3852")
        for nombre in NOMBRES_DE_LOS_PASOS:
            self.assertEqual(1, primera[nombre])
            self.assertEqual(0, segunda[nombre])
        self.assertEqual(0, primera[COLUMNA_DE_LA_LLAMADA])
        self.assertEqual(1, segunda[COLUMNA_DE_LA_LLAMADA])

    def test_los_seis_pasos_completos_no_inventan_un_estado_que_nadie_dijo(self):
        """`DECISIONES.md` P-1: «completa» no es un valor que nadie haya decidido.

        Lo que consta es que alguien la miró y que los seis salieron: eso lo dicen
        las seis columnas y la firma. Escribir un estado que el dueño no ha nombrado
        sería inventarlo (regla permanente 1).
        """
        def contestar(hoja):
            fila = self.fila_de(hoja, "055-1111-3851")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=fila, column=indice_de(nombre)).value = "Sí"

        self.editar(contestar)
        reconciliar_excel(self.conexion, self.resultado.ruta_del_excel, self.companero_id)

        persona = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3851")
        self.assertIsNone(persona["estado_propuesto"])
        self.assertIsNotNone(persona["propuesto_por"], "pero SÍ consta quién la miró")

    def test_un_paso_que_no_se_traduce_en_incompleta(self):
        def contestar(hoja):
            fila = self.fila_de(hoja, "055-1111-3851")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=fila, column=indice_de(nombre)).value = "Sí"
            hoja.cell(row=fila, column=indice_de("paso_entrevistas")).value = "No"

        self.editar(contestar)
        reconciliar_excel(self.conexion, self.resultado.ruta_del_excel, self.companero_id)

        persona = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3851")
        self.assertEqual("incompleta", persona["estado_propuesto"])

    def test_una_respuesta_fuera_del_menu_se_lee_igual_y_se_avisa_con_su_fila(self):
        """«El menú es una ayuda, no una reja», y el aviso lleva el número de fila."""
        def contestar(hoja):
            fila = self.fila_de(hoja, "055-1111-3851")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=fila, column=indice_de(nombre)).value = "SI"
            hoja.cell(row=fila, column=indice_de("paso_cita_del_templo")).value = "regular"
            otra = self.fila_de(hoja, "055-1111-3852")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=otra, column=indice_de(nombre)).value = "no"

        self.editar(contestar)
        vuelta = reconciliar_excel(
            self.conexion, self.resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(1, len(vuelta.descartadas))
        descartada = vuelta.descartadas[0]
        self.assertEqual(PRIMERA_FILA_DE_DATOS, descartada.numero)
        self.assertIn("3. Cita del templo", descartada.motivo)
        self.assertIn("regular", descartada.motivo)
        self.assertEqual(1, len(vuelta.aplicadas), "«no» en minúscula SÍ se entiende")

    def test_una_fila_en_blanco_no_es_un_no(self):
        """Sin contestar deja a la persona a medias, no reprobada."""
        def contestar(hoja):
            fila = self.fila_de(hoja, "055-1111-3851")
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=fila, column=indice_de(nombre)).value = "Sí"

        self.editar(contestar)
        vuelta = reconciliar_excel(
            self.conexion, self.resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual((PRIMERA_FILA_DE_DATOS + 1,), vuelta.sin_nada_que_proponer)
        sin_tocar = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3852")
        for nombre in NOMBRES_DE_LOS_PASOS:
            self.assertIsNone(
                sin_tocar[nombre], "una celda en blanco NO se guarda como un «no»"
            )
        self.assertIsNone(sin_tocar["propuesto_por"])

    def test_borrar_la_clave_descarta_la_fila_y_no_la_busca_por_nombre(self):
        def borrar(hoja):
            fila = self.fila_de(hoja, "055-1111-3851")
            hoja.cell(row=fila, column=indice_de("clave")).value = None
            for nombre in NOMBRES_DE_LOS_PASOS:
                hoja.cell(row=fila, column=indice_de(nombre)).value = "No"

        self.editar(borrar)
        vuelta = reconciliar_excel(
            self.conexion, self.resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(1, len(vuelta.descartadas))
        self.assertIn("clave", vuelta.descartadas[0].motivo)
        self.assertIn("nombre", vuelta.descartadas[0].motivo.lower())
        persona = resolver_persona_por_par(self.conexion, "CASP2610", "055-1111-3851")
        self.assertIsNone(persona["propuesto_por"])

    def test_la_vuelta_no_marca_nada_como_verificado(self):
        """Regla permanente 5: lo que trae el compañero es una PROPUESTA."""
        def contestar(hoja):
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
                for nombre in NOMBRES_DE_LOS_PASOS:
                    hoja.cell(row=fila, column=indice_de(nombre)).value = "Sí"

        self.editar(contestar)
        reconciliar_excel(self.conexion, self.resultado.ruta_del_excel, self.companero_id)

        verificados = self.conexion.execute(
            "SELECT COUNT(*) FROM procedencia_campo WHERE verificado = 1"
        ).fetchone()[0]
        self.assertEqual(0, verificados)


if __name__ == "__main__":
    unittest.main()
