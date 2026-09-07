"""La persona sin MRN se avisa al generar, y lo descartado de la vuelta se guarda.

**El hallazgo ALTO de la auditoría final de QA (2026-09-03), con su medición.** Ida
y vuelta con seis perfiles: los cinco con MRN volvieron con 7/7 pasos y **la que no
tiene MRN volvió con 0/7**. Y no es un caso raro de laboratorio: el PDF real del
dueño trae **1 de 4** personas sin MRN.

El circuito es correcto pieza a pieza y falla como conjunto. La vuelta casa por
`numero_caso` + `mrn` y **nunca por nombre** (`DECISIONES.md`, 2026-09-02): un
acento de más crea un registro fantasma, así que descartar la fila es lo que hay
que hacer. Lo que faltaba son las dos mitades que convierten un descarte correcto
en algo que alguien puede arreglar:

  1. **Avisar al GENERAR**, nombrando a quien no tiene cédula. Es el único momento
     en que todavía se puede arreglar sin gastar el trabajo de nadie: después, el
     compañero ya contestó y ese trabajo se pierde entero.
  2. **Guardar lo descartado.** Hasta hoy la lista se veía en una ventana y se
     perdía al cerrarla, cosa que `interfaz/descartados.py` decía de sí mismo. Con
     cientos de filas, lo que no queda escrito no existe.

**Lo que estas pruebas NO cubren, y es del dueño:** teclear a mano los pasos de la
persona sin MRN cuando el Excel vuelve. Eso exige decidir quién firma esa
corrección y nadie lo ha decidido.
"""

import shutil
import tempfile
import unittest
from pathlib import Path

from datos.asignaciones import asignar_caso
from datos.companeros import alta_de_companero
from datos.descartadas import contar_filas_descartadas, filas_descartadas
from datos.repositorio import alta_de_caso, alta_de_persona
from paquete.exportacion import exportar_paquete
from paquete.reconciliacion import reconciliar_excel
from pruebas.comun import PruebaConBaseTemporal

NOMBRE_SIN_MRN = "PEREZ RAMIREZ, ANA LUCIA"
NOMBRE_CON_MRN = "ANONIMO, JOSE MIGUEL"


class ElPaqueteAvisaDeQuienNoTieneCedula(PruebaConBaseTemporal):
    """Dado un caso con una persona sin MRN, cuando se genera el paquete del
    compañero, entonces el aviso la nombra y dice qué le va a pasar."""

    def setUp(self):
        super().setUp()
        self.carpeta_del_paquete = Path(tempfile.mkdtemp(prefix="fichas_paquete_"))
        self.addCleanup(shutil.rmtree, self.carpeta_del_paquete, ignore_errors=True)
        self.companero_id = alta_de_companero(self.conexion, "Ana")
        self.caso_id = alta_de_caso(
            self.conexion,
            numero_caso="CASP2609",
            unidad_numero="123456",
            fecha_viaje="2026-09-08",
            captura_manual=0,
            ruta_pdf=None,
            pagina_pdf=1,
            unidad_nombre="Cuatricentenaria",
        ).id
        asignar_caso(self.conexion, self.caso_id, self.companero_id)

    def _persona(self, nombre, mrn, fila):
        alta_de_persona(
            self.conexion, self.caso_id, mrn=mrn, nombre=nombre,
            fila_formulario=fila, pagina_pdf=1,
        )

    def _generar(self):
        return exportar_paquete(
            self.conexion, self.companero_id, self.carpeta_del_paquete
        )

    def test_el_aviso_nombra_a_la_persona_sin_mrn(self):
        """«Hay 1 persona sin MRN» obliga a abrir el Excel a buscarla."""
        self._persona(NOMBRE_CON_MRN, "055-1111-3853", 1)
        self._persona(NOMBRE_SIN_MRN, None, 2)
        avisos = " ".join(self._generar().avisos)
        self.assertIn(NOMBRE_SIN_MRN, avisos)

    def test_el_aviso_dice_que_esa_persona_no_podra_recibir_respuesta(self):
        """El hecho primero y la consecuencia después, como el resto de los avisos."""
        self._persona(NOMBRE_SIN_MRN, None, 1)
        avisos = " ".join(self._generar().avisos)
        self.assertIn("cédula", avisos)
        self.assertIn("NUNCA por nombre", avisos)

    def test_sin_nadie_sin_mrn_no_sale_ese_aviso(self):
        """Un aviso que sale siempre no lo lee nadie, y entonces el que importa tampoco."""
        self._persona(NOMBRE_CON_MRN, "055-1111-3853", 1)
        avisos = " ".join(self._generar().avisos)
        self.assertNotIn("cédula de miembro", avisos)

    def test_el_paquete_se_genera_igual_con_su_fila_dentro(self):
        """Avisar no es negarse: esa persona va en la hoja como las demás."""
        self._persona(NOMBRE_CON_MRN, "055-1111-3853", 1)
        self._persona(NOMBRE_SIN_MRN, None, 2)
        resultado = self._generar()
        self.assertTrue(resultado.ruta_del_excel.is_file())
        self.assertEqual(len(resultado.casos), 1)

    def test_dos_personas_sin_mrn_salen_las_dos_nombradas(self):
        self._persona(NOMBRE_SIN_MRN, None, 1)
        self._persona("SIN CEDULA, OTRO", None, 2)
        avisos = " ".join(self._generar().avisos)
        self.assertIn(NOMBRE_SIN_MRN, avisos)
        self.assertIn("SIN CEDULA, OTRO", avisos)


class LoDescartadoDeLaVueltaSeQueda(PruebaConBaseTemporal):
    """Dado un Excel devuelto con una fila que no casa con nadie, cuando se carga,
    entonces su renglón se guarda y se puede volver a mirar mañana."""

    def setUp(self):
        super().setUp()
        self.carpeta_del_paquete = Path(tempfile.mkdtemp(prefix="fichas_vuelta_"))
        self.addCleanup(shutil.rmtree, self.carpeta_del_paquete, ignore_errors=True)
        self.companero_id = alta_de_companero(self.conexion, "Ana")
        self.caso_id = alta_de_caso(
            self.conexion,
            numero_caso="CASP2609",
            unidad_numero="123456",
            fecha_viaje="2026-09-08",
            captura_manual=0,
            ruta_pdf=None,
            pagina_pdf=1,
            unidad_nombre="Cuatricentenaria",
        ).id
        asignar_caso(self.conexion, self.caso_id, self.companero_id)
        alta_de_persona(
            self.conexion, self.caso_id, mrn="055-1111-3853", nombre=NOMBRE_CON_MRN,
            fila_formulario=1, pagina_pdf=1,
        )
        alta_de_persona(
            self.conexion, self.caso_id, mrn=None, nombre=NOMBRE_SIN_MRN,
            fila_formulario=2, pagina_pdf=1,
        )
        self.paquete = exportar_paquete(
            self.conexion, self.companero_id, self.carpeta_del_paquete
        )
        self._contestar_todo_que_si()

    def _contestar_todo_que_si(self):
        """Rellena las siete respuestas de las dos filas, como haría el compañero."""
        from openpyxl import load_workbook

        from paquete.columnas import COLUMNAS, FILA_DE_LA_CABECERA, PRIMERA_FILA_DE_DATOS
        from datos.pasos import COLUMNA_DE_LA_LLAMADA, PASOS

        libro = load_workbook(str(self.paquete.ruta_del_excel))
        hoja = libro.active
        a_rellenar = {nombre for nombre, _ in PASOS} | {COLUMNA_DE_LA_LLAMADA}
        columnas = {
            columna.nombre: numero
            for numero, columna in enumerate(COLUMNAS, start=1)
        }
        for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
            for nombre in a_rellenar:
                hoja.cell(row=fila, column=columnas[nombre], value="Sí")
        libro.save(str(self.paquete.ruta_del_excel))
        self.assertGreater(hoja.max_row, FILA_DE_LA_CABECERA)

    def test_la_fila_sin_mrn_se_descarta_y_deja_su_renglon(self):
        vuelta = reconciliar_excel(
            self.conexion, self.paquete.ruta_del_excel, self.companero_id
        )
        self.assertEqual(len(vuelta.descartadas), 1)
        self.assertEqual(contar_filas_descartadas(self.conexion), 1)

    def test_el_renglon_guardado_dice_de_quien_era_y_por_que_no_entro(self):
        """Sin el motivo, el renglón dice «algo falló» y no se puede hacer nada."""
        reconciliar_excel(
            self.conexion, self.paquete.ruta_del_excel, self.companero_id
        )
        renglon = filas_descartadas(self.conexion)[0]
        self.assertEqual(renglon["companero"], "Ana")
        self.assertTrue(renglon["motivo"])
        self.assertIsNotNone(renglon["fila_excel"])

    def test_el_renglon_guardado_dice_de_que_archivo_salio_esa_fila(self):
        """«Fila 8» no lleva a ningún sitio si hay tres Excel en la carpeta."""
        reconciliar_excel(
            self.conexion, self.paquete.ruta_del_excel, self.companero_id
        )
        renglon = filas_descartadas(self.conexion)[0]
        self.assertEqual(renglon["ruta_excel"], str(self.paquete.ruta_del_excel))

    def test_el_renglon_sobrevive_a_cerrar_el_programa(self):
        """La prueba del criterio: se abre otra conexión sobre el mismo archivo.

        Es lo que la ventana NO hacía: se veía mientras estaba abierta y al
        cerrarla se perdía, y con ella la única pista de que un compañero había
        hecho un trabajo que nadie recogió.
        """
        from datos.conexion import abrir_conexion

        reconciliar_excel(
            self.conexion, self.paquete.ruta_del_excel, self.companero_id
        )
        self.conexion.close()
        otra = abrir_conexion(self.ruta_de_la_base)
        try:
            self.assertEqual(contar_filas_descartadas(otra), 1)
        finally:
            otra.close()
        self.conexion = abrir_conexion(self.ruta_de_la_base)

    def test_la_fila_que_si_casa_se_aplica_igual(self):
        """Guardar el descarte no puede cambiar lo que sí entra."""
        vuelta = reconciliar_excel(
            self.conexion, self.paquete.ruta_del_excel, self.companero_id
        )
        self.assertEqual(len(vuelta.aplicadas), 1)

    def test_una_vuelta_sin_descartes_no_deja_ningun_renglon(self):
        """La otra mitad: una tabla que se llena siempre deja de informar."""
        self.conexion.execute("DELETE FROM personas WHERE mrn IS NULL")
        paquete = exportar_paquete(
            self.conexion, self.companero_id, self.carpeta_del_paquete
        )
        self.paquete = paquete
        self._contestar_todo_que_si()
        reconciliar_excel(self.conexion, paquete.ruta_del_excel, self.companero_id)
        self.assertEqual(contar_filas_descartadas(self.conexion), 0)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
