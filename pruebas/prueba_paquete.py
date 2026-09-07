"""FASE 6: companeros, asignaciones, paquete y la vuelta del Excel.

Cada prueba de aqui sale de un criterio de aceptacion de `PENDIENTES.md`, y el
criterio esta citado en el nombre o en la primera linea. **No sale de mirar el
codigo que acabo de escribir**: una prueba escrita mirando el codigo solo confirma
lo que su autor entendio, y pasa en verde estando mal.

El criterio de cierre del pase se prueba entero en `PruebaDelCriterioDeCierre`:
exportar 3 casos, editar el Excel POR FUERA cambiando nombres y borrando una fila,
volver a cargarlo, y comprobar que cada estado cae en la persona correcta.

Ninguna prueba toca la base real de Miguel: todas trabajan sobre `PruebaConBaseTemporal`.
"""

import unittest

from openpyxl import load_workbook

from datos.asignaciones import (
    asignar_caso,
    asignar_casos,
    casos_asignados,
    companeros_del_caso,
    retirar_caso,
)
from datos.companeros import (
    alta_de_companero,
    companeros_activos,
    desactivar_companero,
    leer_companero,
    reactivar_companero,
    renombrar_companero,
)
from datos.procedencia import (
    guardar_procedencia_de_campo,
    leer_procedencia_de_campo,
    marcar_campo_verificado,
)
from datos.propuestas import propuestas_del_caso, resolver_persona_por_par
from datos.repositorio import alta_de_caso, alta_de_persona
from datos.validacion import ErrorDeValidacion
from datos.pasos import COLUMNA_DE_LA_LLAMADA, NOMBRES_DE_LOS_PASOS
from paquete.columnas import (
    FILA_DE_LA_CABECERA,
    NOMBRE_DE_LA_HOJA,
    PRIMERA_FILA_DE_DATOS,
    indice_de,
)
from paquete.exportacion import exportar_paquete, filas_de_trabajo
from paquete.reconciliacion import reconciliar_excel
from pruebas.comun import PruebaConBaseTemporal


def contestar_los_seis_pasos(hoja, fila, respuesta="Sí", llamo=None):
    """Rellena las siete casillas que rellena el companero en esa fila.

    Existe porque desde el 2026-09-03 la hoja **no pregunta un estado**: pregunta
    los seis pasos de «Preparación para las ordenanzas», y el estado se deduce de
    ellos. Una prueba que escribiera en una columna «estado» estaría midiendo un
    modelo que ya no es el del programa.
    """
    for nombre in NOMBRES_DE_LOS_PASOS:
        hoja.cell(row=fila, column=indice_de(nombre)).value = respuesta
    if llamo is not None:
        hoja.cell(row=fila, column=indice_de(COLUMNA_DE_LA_LLAMADA)).value = llamo


def _contar(conexion, tabla_literal):
    """Cuenta las filas de una tabla. Cada consulta es literal, sin interpolar."""
    if tabla_literal == "personas":
        return conexion.execute("SELECT COUNT(*) FROM personas").fetchone()[0]
    if tabla_literal == "companeros":
        return conexion.execute("SELECT COUNT(*) FROM companeros").fetchone()[0]
    if tabla_literal == "casos":
        return conexion.execute("SELECT COUNT(*) FROM casos").fetchone()[0]
    raise AssertionError(f"La prueba no sabe contar la tabla {tabla_literal!r}.")


class BaseConTresCasos(PruebaConBaseTemporal):
    """Tres casos con dos personas cada uno, y un companero al que asignarselos."""

    def setUp(self):
        super().setUp()
        self.companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        self.casos = {}
        for indice, numero in enumerate(("CASP2609", "CBSP2609", "CCSP2609"), start=1):
            resultado = alta_de_caso(
                self.conexion,
                numero,
                unidad_numero="123456",
                unidad_nombre="Paramaribo Branch",
                fecha_viaje="2026-09-08",
                pagina_pdf=indice,
            )
            self.casos[numero] = resultado.id
            for fila in (1, 2):
                alta_de_persona(
                    self.conexion,
                    resultado.id,
                    mrn=f"05{indice}-1111-385{fila}",
                    nombre=f"José Miguel Anonimo {indice}{fila}",
                    fila_formulario=fila,
                    pagina_pdf=indice,
                )
        asignar_casos(self.conexion, list(self.casos.values()), self.companero_id)

    def exportar(self):
        """Genera el paquete en la carpeta temporal y devuelve el resultado."""
        return exportar_paquete(self.conexion, self.companero_id, self.carpeta_temporal)


class PruebaDeCompaneros(PruebaConBaseTemporal):
    """FASE 6, criterios 3 y 4c: se desactivan, no se borran."""

    def test_desactivar_lo_saca_del_desplegable_sin_borrar_la_fila(self):
        """Criterio 4c: `COUNT(*)` devuelve el mismo numero antes y despues."""
        companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        antes = _contar(self.conexion, "companeros")

        self.assertTrue(desactivar_companero(self.conexion, companero_id))

        self.assertEqual(antes, _contar(self.conexion, "companeros"))
        self.assertEqual([], companeros_activos(self.conexion))
        self.assertIsNotNone(leer_companero(self.conexion, companero_id))

    def test_la_verificacion_de_un_desactivado_sigue_ensenando_su_nombre(self):
        """Criterio 3, la mitad que se olvida: el historial no se puede perder."""
        companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        caso_id = alta_de_caso(self.conexion, "CASP2609").id
        guardar_procedencia_de_campo(
            self.conexion, "casos", caso_id, "fecha_viaje", "ocr", confianza=0.9
        )
        marcar_campo_verificado(
            self.conexion, "casos", caso_id, "fecha_viaje", companero_id
        )

        desactivar_companero(self.conexion, companero_id)

        procedencia = leer_procedencia_de_campo(
            self.conexion, "casos", caso_id, "fecha_viaje"
        )
        self.assertEqual(1, procedencia["verificado"])
        firmante = leer_companero(self.conexion, procedencia["verificado_por"])
        self.assertEqual("Ana Pérez", firmante["nombre"])

    def test_desactivar_dos_veces_no_mueve_la_fecha_de_la_primera(self):
        companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        desactivar_companero(self.conexion, companero_id)
        primera = leer_companero(self.conexion, companero_id)["desactivado_en"]

        self.assertFalse(desactivar_companero(self.conexion, companero_id))

        self.assertEqual(
            primera, leer_companero(self.conexion, companero_id)["desactivado_en"]
        )

    def test_reactivar_lo_devuelve_al_desplegable_y_borra_la_fecha(self):
        companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        desactivar_companero(self.conexion, companero_id)

        self.assertTrue(reactivar_companero(self.conexion, companero_id))

        companero = leer_companero(self.conexion, companero_id)
        self.assertEqual(1, companero["activo"])
        self.assertIsNone(companero["desactivado_en"])
        self.assertEqual(1, len(companeros_activos(self.conexion)))

    def test_un_nombre_vacio_no_se_guarda(self):
        with self.assertRaises(ErrorDeValidacion):
            alta_de_companero(self.conexion, "   ")

    def test_dos_companeros_pueden_llamarse_igual(self):
        """`docs/ARQUITECTURA.md` §2.4: `nombre` NO es UNIQUE, a proposito."""
        alta_de_companero(self.conexion, "José Anonimo")
        alta_de_companero(self.conexion, "José Anonimo")
        self.assertEqual(2, len(companeros_activos(self.conexion)))

    def test_renombrar_no_toca_el_estado(self):
        companero_id = alta_de_companero(self.conexion, "Ana Peres")
        renombrar_companero(self.conexion, companero_id, "Ana Pérez")
        companero = leer_companero(self.conexion, companero_id)
        self.assertEqual("Ana Pérez", companero["nombre"])
        self.assertEqual(1, companero["activo"])

    def test_no_hay_ninguna_funcion_de_borrado_en_el_modulo(self):
        """Criterio 4b: la puerta se cierra no ofreciendo la operacion."""
        import datos.companeros as modulo

        publicas = [
            nombre for nombre in dir(modulo) if not nombre.startswith("_")
        ]
        sospechosas = [
            nombre
            for nombre in publicas
            if "borrar" in nombre or "eliminar" in nombre or "delete" in nombre
        ]
        self.assertEqual([], sospechosas)


class PruebaDeAsignaciones(PruebaConBaseTemporal):
    """Las asignaciones se retiran, no se borran (`DECISIONES.md`, P-4)."""

    def setUp(self):
        super().setUp()
        self.companero_id = alta_de_companero(self.conexion, "Ana Pérez")
        self.caso_id = alta_de_caso(self.conexion, "CASP2609").id

    def test_asignar_dos_veces_no_crea_dos_asignaciones_vivas(self):
        primera = asignar_caso(self.conexion, self.caso_id, self.companero_id)
        segunda = asignar_caso(self.conexion, self.caso_id, self.companero_id)
        self.assertEqual(primera, segunda)
        self.assertEqual(1, len(casos_asignados(self.conexion, self.companero_id)))

    def test_retirar_y_volver_a_asignar_funciona(self):
        """El indice unico es PARCIAL: solo prohibe DOS VIVAS a la vez."""
        asignar_caso(self.conexion, self.caso_id, self.companero_id)
        self.assertTrue(retirar_caso(self.conexion, self.caso_id, self.companero_id))
        self.assertEqual([], casos_asignados(self.conexion, self.companero_id))

        asignar_caso(self.conexion, self.caso_id, self.companero_id)

        self.assertEqual(1, len(casos_asignados(self.conexion, self.companero_id)))
        filas = self.conexion.execute("SELECT COUNT(*) FROM asignaciones").fetchone()[0]
        self.assertEqual(2, filas, "la asignacion retirada tiene que seguir en la tabla")

    def test_no_se_le_asigna_trabajo_a_un_companero_desactivado(self):
        desactivar_companero(self.conexion, self.companero_id)
        with self.assertRaises(ErrorDeValidacion):
            asignar_caso(self.conexion, self.caso_id, self.companero_id)

    def test_companeros_del_caso_dice_quien_lo_lleva(self):
        asignar_caso(self.conexion, self.caso_id, self.companero_id)
        llevan = companeros_del_caso(self.conexion, self.caso_id)
        self.assertEqual(["Ana Pérez"], [fila["nombre"] for fila in llevan])


class PruebaDelPaqueteQueSale(BaseConTresCasos):
    """FASE 6, criterio 5: el paquete abre y trae EXACTAMENTE esos casos."""

    def test_el_excel_trae_las_seis_personas_de_los_tres_casos(self):
        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        numeros = {
            hoja.cell(row=fila, column=indice_de("numero_caso")).value
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1)
        }
        libro.close()

        self.assertEqual(set(self.casos), numeros)
        self.assertEqual(3, len(numeros), "ni un caso de más")
        self.assertEqual(6, hoja.max_row - FILA_DE_LA_CABECERA, "dos personas por caso")

    def test_no_entra_un_caso_que_no_esta_asignado(self):
        ajeno = alta_de_caso(self.conexion, "CZSP2609").id
        alta_de_persona(self.conexion, ajeno, mrn="009-9999-9999", nombre="Ajena")

        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        numeros = {
            hoja.cell(row=fila, column=indice_de("numero_caso")).value
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1)
        }
        libro.close()
        self.assertNotIn("CZSP2609", numeros)

    def test_numero_de_caso_y_mrn_quedan_bloqueados_y_el_resto_no(self):
        """El requisito explicito del pase, comprobado sobre el archivo escrito."""
        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        self.assertTrue(hoja.protection.sheet, "sin proteger la hoja, el bloqueo no hace nada")
        for columna in ("numero_caso", "mrn"):
            celda = hoja.cell(row=PRIMERA_FILA_DE_DATOS, column=indice_de(columna))
            self.assertTrue(celda.protection.locked, f"{columna} tenía que estar bloqueada")
        for columna in NOMBRES_DE_LOS_PASOS + (COLUMNA_DE_LA_LLAMADA, "nombre"):
            celda = hoja.cell(row=PRIMERA_FILA_DE_DATOS, column=indice_de(columna))
            self.assertFalse(
                celda.protection.locked, f"{columna} tenía que poder escribirse"
            )
        libro.close()

    def test_las_siete_columnas_del_companero_llevan_menu_de_si_y_no(self):
        """Un menú por cada casilla que rellena, y ninguno más.

        Siete y no una: son los seis pasos más «¿Llamó al líder?». Si faltara uno,
        esa columna se rellenaría a mano y volvería escrita de siete maneras.
        """
        from datos.pasos import RESPUESTAS

        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        validaciones = hoja.data_validations.dataValidation
        formulas = {validacion.formula1 for validacion in validaciones}
        cuantas = len(validaciones)
        libro.close()

        self.assertEqual(len(NOMBRES_DE_LOS_PASOS) + 1, cuantas)
        self.assertEqual({'"' + ",".join(RESPUESTAS) + '"'}, formulas)

    def test_el_mrn_conserva_sus_ceros_de_delante(self):
        """Sin `number_format = '@'` Excel se los come y el MRN deja de casar."""
        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        valores = {
            hoja.cell(row=fila, column=indice_de("mrn")).value
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1)
        }
        libro.close()
        self.assertIn("051-1111-3851", valores)

    def test_un_companero_sin_casos_no_produce_un_paquete_vacio(self):
        otro = alta_de_companero(self.conexion, "Sin Trabajo")
        with self.assertRaises(ErrorDeValidacion):
            exportar_paquete(self.conexion, otro, self.carpeta_temporal)

    def test_el_estado_sale_en_blanco_aunque_haya_una_propuesta_anterior(self):
        """Lo que se manda es lo que hay que mirar, no lo que otro contesto."""
        from datos.propuestas import guardar_propuesta

        persona = resolver_persona_por_par(self.conexion, "CASP2609", "051-1111-3851")
        guardar_propuesta(
            self.conexion, persona["id"], "incompleta", "de la ronda anterior",
            self.companero_id,
        )

        filas = filas_de_trabajo(
            self.conexion, casos_asignados(self.conexion, self.companero_id)
        )

        for fila in filas:
            for nombre in NOMBRES_DE_LOS_PASOS + (COLUMNA_DE_LA_LLAMADA,):
                self.assertIsNone(
                    fila.get(nombre),
                    f"«{nombre}» tenía que salir en blanco: lo que se manda es lo "
                    "que hay que mirar, no lo que otro contestó",
                )


class PruebaDeLaVuelta(BaseConTresCasos):
    """FASE 6, criterios 1, 2 y 6: que actualiza, que se descarta, y la idempotencia."""

    def _editar(self, ruta, cambios):
        """Edita el Excel POR FUERA, como haria el companero. Devuelve la ruta.

        Se abre y se guarda con `openpyxl` y no se reescribe desde cero: eso es lo
        que hace un programa de hoja de calculo, y lo que se quiere probar es que
        la vuelta aguanta un archivo que paso por otras manos.
        """
        libro = load_workbook(str(ruta))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        cambios(hoja)
        libro.save(str(ruta))
        libro.close()
        return ruta

    def _fila_de(self, hoja, numero_caso, mrn):
        for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
            if (
                hoja.cell(row=fila, column=indice_de("numero_caso")).value == numero_caso
                and hoja.cell(row=fila, column=indice_de("mrn")).value == mrn
            ):
                return fila
        raise AssertionError(f"No hay fila para {numero_caso} + {mrn}.")

    def test_criterio_1_un_acento_distinto_actualiza_y_no_crea_nada(self):
        """El nombre cambia, el par no: la fila ACTUALIZA. Conteo igual."""
        resultado = self.exportar()
        antes = _contar(self.conexion, "personas")

        def cambiar(hoja):
            fila = self._fila_de(hoja, "CASP2609", "051-1111-3851")
            hoja.cell(row=fila, column=indice_de("nombre")).value = "Jose Miguél Órtiz 11"
            contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, cambiar)
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(antes, _contar(self.conexion, "personas"))
        self.assertEqual(1, len(vuelta.aplicadas))
        persona = resolver_persona_por_par(self.conexion, "CASP2609", "051-1111-3851")
        self.assertEqual("incompleta", persona["estado_propuesto"])
        self.assertEqual(
            "José Miguel Anonimo 11",
            persona["nombre"],
            "el nombre de la base NO se toca: la vuelta no escribe nombres",
        )

    def test_criterio_2_un_par_que_no_existe_no_se_inserta_y_se_descarta(self):
        resultado = self.exportar()
        antes_personas = _contar(self.conexion, "personas")
        antes_casos = _contar(self.conexion, "casos")

        def inventar(hoja):
            fila = hoja.max_row + 1
            hoja.cell(row=fila, column=indice_de("numero_caso")).value = "CXSP2609"
            hoja.cell(row=fila, column=indice_de("mrn")).value = "007-7777-7777"
            hoja.cell(row=fila, column=indice_de("nombre")).value = "Persona Inventada"
            hoja.cell(row=fila, column=indice_de("clave")).value = "CXSP2609:007-7777-7777"
            contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, inventar)
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(antes_personas, _contar(self.conexion, "personas"))
        self.assertEqual(antes_casos, _contar(self.conexion, "casos"))
        self.assertEqual(1, len(vuelta.descartadas))
        descartada = vuelta.descartadas[0]
        self.assertEqual("CXSP2609", descartada.numero_caso)
        self.assertIn("no existe en la base", descartada.motivo)

    def test_criterio_6_reconciliar_dos_veces_no_duplica_nada(self):
        resultado = self.exportar()

        def marcar(hoja):
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
                contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, marcar)
        reconciliar_excel(self.conexion, resultado.ruta_del_excel, self.companero_id)
        despues_de_la_primera = _contar(self.conexion, "personas")

        segunda = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(despues_de_la_primera, _contar(self.conexion, "personas"))
        self.assertEqual(6, len(segunda.aplicadas))
        self.assertEqual(0, len(segunda.descartadas))

    def test_una_fila_sin_mrn_se_descarta_y_no_se_empareja_por_nombre(self):
        """La regla que sostiene todo: sin el par, NUNCA se busca por nombre."""
        resultado = self.exportar()

        def borrar_el_mrn(hoja):
            fila = self._fila_de(hoja, "CASP2609", "051-1111-3851")
            hoja.cell(row=fila, column=indice_de("mrn")).value = None
            hoja.cell(row=fila, column=indice_de("clave")).value = None
            contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, borrar_el_mrn)
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(1, len(vuelta.descartadas))
        self.assertIn("nombre", vuelta.descartadas[0].motivo.lower())
        persona = resolver_persona_por_par(self.conexion, "CASP2609", "051-1111-3851")
        self.assertIsNone(persona["estado_propuesto"])

    def test_una_respuesta_fuera_del_menu_se_lee_igual_y_se_avisa(self):
        """«El menú es una ayuda, no una reja» — proyecto viejo, `salida/asignacion.py`.

        Lo que se comprueba son las dos mitades: que el archivo se lee entero
        aunque una celda traiga algo escrito a mano, y que esa fila se avisa con su
        número y con el rótulo de la columna. Un archivo rechazado entero por una
        celda rara le costaría al compañero la ronda completa.
        """
        resultado = self.exportar()

        def contestar_a_mano(hoja):
            fila = self._fila_de(hoja, "CASP2609", "051-1111-3851")
            contestar_los_seis_pasos(hoja, fila, "Sí")
            hoja.cell(row=fila, column=indice_de("paso_entrevistas")).value = "más o menos"
            otra = self._fila_de(hoja, "CBSP2609", "052-1111-3851")
            contestar_los_seis_pasos(hoja, otra, "Sí")

        self._editar(resultado.ruta_del_excel, contestar_a_mano)
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(1, len(vuelta.descartadas))
        motivo = vuelta.descartadas[0].motivo
        self.assertIn("5. Entrevistas", motivo)
        self.assertIn("más o menos", motivo)
        self.assertEqual(
            1, len(vuelta.aplicadas), "la fila buena del resto del archivo sí entra"
        )
        sin_aplicar = resolver_persona_por_par(
            self.conexion, "CASP2609", "051-1111-3851"
        )
        self.assertIsNone(
            sin_aplicar["propuesto_por"],
            "no se aplica media fila: los cinco pasos legibles tampoco entran",
        )

    def test_el_mismo_par_repetido_en_el_archivo_solo_entra_una_vez(self):
        resultado = self.exportar()

        def repetir(hoja):
            origen = self._fila_de(hoja, "CASP2609", "051-1111-3851")
            contestar_los_seis_pasos(hoja, origen, "No")
            destino = hoja.max_row + 1
            hoja.cell(row=destino, column=indice_de("numero_caso")).value = "CASP2609"
            hoja.cell(row=destino, column=indice_de("mrn")).value = "051-1111-3851"
            # La clave se copia entera de la fila original, con su id de caso: una
            # clave a medias caeria en «ambigua» y esta prueba mide la REPETIDA.
            hoja.cell(row=destino, column=indice_de("clave")).value = hoja.cell(
                row=origen, column=indice_de("clave")
            ).value
            contestar_los_seis_pasos(hoja, destino, "Sí")

        self._editar(resultado.ruta_del_excel, repetir)
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        self.assertEqual(1, len(vuelta.aplicadas))
        self.assertEqual(1, len(vuelta.descartadas))
        self.assertIn("ya venía en la fila", vuelta.descartadas[0].motivo)
        persona = resolver_persona_por_par(self.conexion, "CASP2609", "051-1111-3851")
        self.assertEqual("incompleta", persona["estado_propuesto"])

    def test_la_vuelta_no_marca_nada_como_verificado(self):
        """Regla permanente 5: lo que trae el companero es una PROPUESTA."""
        resultado = self.exportar()

        def marcar(hoja):
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
                contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, marcar)
        reconciliar_excel(self.conexion, resultado.ruta_del_excel, self.companero_id)

        verificados = self.conexion.execute(
            "SELECT COUNT(*) FROM procedencia_campo WHERE verificado = 1"
        ).fetchone()[0]
        self.assertEqual(0, verificados)

    def test_la_vuelta_SI_marca_el_estado_del_caso_y_dice_quien_lo_marco(self):
        """⚠️ **Esta prueba dio la vuelta el 2026-09-03, por decisión del dueño.**

        Hasta ese día exigía lo contrario —«el estado del caso lo escribe Miguel, no
        un archivo que llega»— y él lo cambió con estas palabras: *«el documento que
        ellos llenan de Excel es el que marca»*. La carga escribe el estado, y queda
        firmado con el compañero que lo devolvió, no en blanco.

        **La regla permanente 5 NO se toca, y son dos cosas distintas.** Lo que la
        vuelta escribe es `casos.estado_recomendacion` —el ESTADO del documento—;
        lo que sigue exigiendo que Miguel pulse el botón es
        `procedencia_campo.verificado`, la FIRMA de cada campo leído. Que sigue
        intacto lo mide la prueba de al lado,
        `test_la_vuelta_no_marca_nada_como_verificado`, que no cambió ni una línea.
        """
        resultado = self.exportar()

        def marcar(hoja):
            fila = self._fila_de(hoja, "CASP2609", "051-1111-3851")
            contestar_los_seis_pasos(hoja, fila, "No")

        self._editar(resultado.ruta_del_excel, marcar)
        reconciliar_excel(self.conexion, resultado.ruta_del_excel, self.companero_id)

        fila = self.conexion.execute(
            "SELECT estado_recomendacion, estado_marcado_por, estado_marcado_origen "
            "FROM casos WHERE numero_caso = ?",
            ("CASP2609",),
        ).fetchone()
        self.assertIsNotNone(fila["estado_recomendacion"])
        self.assertEqual(fila["estado_marcado_por"], self.companero_id)
        self.assertIn("por_verificar.xlsx", fila["estado_marcado_origen"])


class PruebaDelCriterioDeCierre(BaseConTresCasos):
    """El criterio de cierre del pase, entero y en una sola prueba.

    «Se exporta un paquete de 3 casos; se edita el Excel por fuera cambiando
    nombres a propósito y borrando una fila; se vuelve a cargar; y todos los
    estados caen en la persona correcta, con la fila borrada y las inventadas en la
    lista de descartados.»
    """

    def test_el_ciclo_completo(self):
        resultado = self.exportar()
        self.assertEqual(3, len(resultado.casos))
        personas_antes = _contar(self.conexion, "personas")
        casos_antes = _contar(self.conexion, "casos")

        # --- se edita el Excel POR FUERA, como haria el companero -------------
        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]

        esperados = {}
        for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1):
            numero = hoja.cell(row=fila, column=indice_de("numero_caso")).value
            mrn = hoja.cell(row=fila, column=indice_de("mrn")).value
            # Un estado distinto por fila: si dos cayeran en la persona equivocada
            # y se intercambiaran, un estado igual para todas no lo notaria.
            respuesta = "No" if fila % 2 == 0 else "Sí"
            contestar_los_seis_pasos(hoja, fila, respuesta, llamo=respuesta)
            # Los nombres se estropean a proposito, todos.
            hoja.cell(row=fila, column=indice_de("nombre")).value = f"NOMBRE CAMBIADO {fila}"
            esperados[(numero, mrn)] = respuesta

        # Se borra una fila entera, la ultima.
        borrada = (
            hoja.cell(row=hoja.max_row, column=indice_de("numero_caso")).value,
            hoja.cell(row=hoja.max_row, column=indice_de("mrn")).value,
        )
        hoja.delete_rows(hoja.max_row)
        esperados.pop(borrada)

        # Y se inventan dos filas que no existen en la base.
        for desplazamiento, (numero, mrn) in enumerate(
            (("CXSP2609", "007-7777-7777"), ("CASP2609", "009-9999-9999"))
        ):
            fila = hoja.max_row + 1 + desplazamiento
            hoja.cell(row=fila, column=indice_de("numero_caso")).value = numero
            hoja.cell(row=fila, column=indice_de("mrn")).value = mrn
            hoja.cell(row=fila, column=indice_de("nombre")).value = "Inventada"
            hoja.cell(row=fila, column=indice_de("clave")).value = f"{numero}:{mrn}"
            contestar_los_seis_pasos(hoja, fila, "No")
        libro.save(str(resultado.ruta_del_excel))
        libro.close()

        # --- se vuelve a cargar ----------------------------------------------
        vuelta = reconciliar_excel(
            self.conexion, resultado.ruta_del_excel, self.companero_id
        )

        # 1. Nada se inserto: los dos conteos son el mismo numero.
        self.assertEqual(personas_antes, _contar(self.conexion, "personas"))
        self.assertEqual(casos_antes, _contar(self.conexion, "casos"))

        # 2. Cada respuesta cayo en SU persona, con los nombres estropeados.
        self.assertEqual(5, len(vuelta.aplicadas))
        for (numero, mrn), respuesta in esperados.items():
            persona = resolver_persona_por_par(self.conexion, numero, mrn)
            self.assertIsNotNone(persona, f"{numero} + {mrn} tenía que resolver")
            esperado = 1 if respuesta == "Sí" else 0
            for nombre in NOMBRES_DE_LOS_PASOS + (COLUMNA_DE_LA_LLAMADA,):
                self.assertEqual(
                    esperado,
                    persona[nombre],
                    f"«{nombre}» de {numero}+{mrn} cayó en la persona equivocada",
                )

        # 3. La fila borrada no quedo tocada: sigue sin propuesta.
        persona_borrada = resolver_persona_por_par(self.conexion, *borrada)
        self.assertIsNone(persona_borrada["estado_propuesto"])
        self.assertIsNone(persona_borrada["propuesto_por"])
        for nombre in NOMBRES_DE_LOS_PASOS:
            self.assertIsNone(persona_borrada[nombre])

        # 4. Las dos inventadas estan en la lista de descartados con su motivo.
        self.assertEqual(2, len(vuelta.descartadas))
        pares_descartados = {
            (descartada.numero_caso, descartada.mrn) for descartada in vuelta.descartadas
        }
        self.assertEqual(
            {("CXSP2609", "007-7777-7777"), ("CASP2609", "009-9999-9999")},
            pares_descartados,
        )
        for descartada in vuelta.descartadas:
            self.assertTrue(descartada.motivo.strip(), "cada descarte lleva su motivo")

        # 5. Y la propuesta lleva el nombre del companero, sin verificar nada.
        propuestas = propuestas_del_caso(self.conexion, self.casos["CASP2609"])
        self.assertTrue(propuestas)
        self.assertEqual("Ana Pérez", propuestas[0]["nombre_del_companero"])
        verificados = self.conexion.execute(
            "SELECT COUNT(*) FROM procedencia_campo WHERE verificado = 1"
        ).fetchone()[0]
        self.assertEqual(0, verificados)


class PruebaDelPdfRecortado(BaseConTresCasos):
    """El PDF que se entrega lleva solo las hojas del caso, no el documento entero."""

    def _pdf_de_tres_paginas(self):
        from pypdf import PdfWriter

        ruta = self.carpeta_temporal / "escaneo.pdf"
        escritor = PdfWriter()
        for _ in range(3):
            escritor.add_blank_page(width=612, height=792)
        with open(ruta, "wb") as archivo:
            escritor.write(archivo)
        return ruta

    def test_solo_van_las_paginas_del_caso(self):
        from pypdf import PdfReader

        ruta_pdf = self._pdf_de_tres_paginas()
        self.conexion.execute(
            "UPDATE casos SET ruta_pdf = ? WHERE numero_caso = ?",
            (str(ruta_pdf), "CASP2609"),
        )

        resultado = self.exportar()

        # El nombre lleva el id del caso pegado desde el 2026-09-03: dos casos
        # pueden compartir numero y el segundo pisaria el archivo del primero.
        esperado = f"CASP2609-{self.casos['CASP2609']}.pdf"
        entregado = [ruta for ruta in resultado.pdf_escritos if ruta.name == esperado]
        self.assertEqual(1, len(entregado), "el caso con PDF tenía que entregar el suyo")
        self.assertEqual(
            1,
            len(PdfReader(str(entregado[0])).pages),
            "el caso vive en la página 1: no se pueden entregar las tres",
        )

    def test_sin_pagina_conocida_NO_se_entrega_el_documento_entero(self):
        """El lado seguro del error: entregar de más enseña casos ajenos."""
        ruta_pdf = self._pdf_de_tres_paginas()
        self.conexion.execute(
            "UPDATE casos SET ruta_pdf = ?, pagina_pdf = NULL WHERE numero_caso = ?",
            (str(ruta_pdf), "CASP2609"),
        )
        self.conexion.execute(
            "UPDATE personas SET pagina_pdf = NULL WHERE caso_id = ?",
            (self.casos["CASP2609"],),
        )

        resultado = self.exportar()

        self.assertEqual([], [r for r in resultado.pdf_escritos if r.name == "CASP2609.pdf"])
        self.assertTrue(
            any("no se sabe qué hojas" in aviso for aviso in resultado.avisos),
            f"tenía que avisar por qué no va el PDF. Avisos: {resultado.avisos}",
        )

    def test_el_caso_va_en_el_excel_aunque_su_pdf_no_exista(self):
        self.conexion.execute(
            "UPDATE casos SET ruta_pdf = ? WHERE numero_caso = ?",
            (str(self.carpeta_temporal / "no_existe.pdf"), "CASP2609"),
        )

        resultado = self.exportar()

        libro = load_workbook(str(resultado.ruta_del_excel))
        hoja = libro[NOMBRE_DE_LA_HOJA]
        numeros = {
            hoja.cell(row=fila, column=indice_de("numero_caso")).value
            for fila in range(PRIMERA_FILA_DE_DATOS, hoja.max_row + 1)
        }
        libro.close()
        self.assertIn("CASP2609", numeros)
        self.assertTrue(any("ya no está" in aviso for aviso in resultado.avisos))


class PruebaDelImportadorLibre(BaseConTresCasos):
    """El Excel que no genero el programa: Miguel dice qué columna es cada cosa."""

    def _hoja_con_titulos_raros(self):
        from openpyxl import Workbook

        ruta = self.carpeta_temporal / "lista_de_ana.xlsx"
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Hoja1"
        # Sin columna «clave»: es un archivo que armo Ana por su cuenta. Ese es el
        # caso que este importador existe para cubrir, y el que obliga a que la
        # reconciliacion sepa caer al par suelto de numero de caso + MRN.
        hoja.append(
            ["Caso", "N. de Registro", "Nombre completo", "Entrevistas", "Listo"]
        )
        hoja.append(["CASP2609", "051-1111-3851", "quien sea", "Sí", "No"])
        hoja.append(["CBSP2609", "052-1111-3852", "otra", "Sí", "Sí"])
        hoja.append(["CZZZ9999", "000-0000-0000", "fantasma", "No", "No"])
        libro.save(str(ruta))
        libro.close()
        return ruta

    def test_propone_la_correspondencia_y_reconcilia_con_el_mismo_motor(self):
        from paquete.mapeo import abrir_para_mapear, campos_sin_asignar, traducir
        from paquete.reconciliacion import reconciliar_filas

        libro, propuesta = abrir_para_mapear(self._hoja_con_titulos_raros())

        self.assertEqual("Caso", propuesta["numero_caso"])
        self.assertEqual("N. de Registro", propuesta["mrn"])
        self.assertEqual("Entrevistas", propuesta["paso_entrevistas"])
        self.assertEqual("Listo", propuesta["paso_listo_para_el_templo"])
        self.assertEqual((), campos_sin_asignar(propuesta))

        antes = _contar(self.conexion, "personas")
        titulos, filas = traducir(libro, propuesta)
        vuelta = reconciliar_filas(self.conexion, titulos, filas, self.companero_id)

        self.assertEqual(antes, _contar(self.conexion, "personas"), "NUNCA inserta")
        self.assertEqual(2, len(vuelta.aplicadas))
        self.assertEqual(1, len(vuelta.descartadas))
        self.assertEqual("CZZZ9999", vuelta.descartadas[0].numero_caso)

    def test_un_titulo_que_no_reconoce_se_queda_sin_asignar_y_no_se_adivina(self):
        from openpyxl import Workbook

        from paquete.mapeo import abrir_para_mapear, campos_sin_asignar

        ruta = self.carpeta_temporal / "rara.xlsx"
        libro_nuevo = Workbook()
        libro_nuevo.active.append(["columna A", "columna B"])
        libro_nuevo.save(str(ruta))
        libro_nuevo.close()

        _, propuesta = abrir_para_mapear(ruta)

        self.assertEqual({}, propuesta)
        self.assertEqual(("numero_caso", "mrn"), campos_sin_asignar(propuesta))


if __name__ == "__main__":
    unittest.main()
