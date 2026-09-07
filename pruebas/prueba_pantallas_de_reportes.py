"""Las dos pantallas de la FASE 8 se abren, se leen y hacen lo que dicen.

Miran la VENTANA, no el valor de retorno de una funcion, por lo mismo que
`pruebas/prueba_aplicacion.py`: lo unico que Miguel puede leer es lo que esta
dibujado. Una metrica que se calcula bien y no se pinta es una metrica que no
existe.

Los cuadros de dialogo se sustituyen durante la prueba. No es un atajo: `informar`
bloquea esperando un clic, y una prueba que abre un cuadro de dialogo se queda
colgada para siempre en la maquina de otro.
"""

import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from datos.archivo import NO_PUDO_VIAJAR, esta_archivado
from datos.repositorio import alta_de_caso, alta_de_persona
from datos.validacion import ErrorDeValidacion
from interfaz.reportes import PantallaDeReportes


def _hay_ventanas():
    """Si no se puede crear un Tk, estas pruebas se omiten en vez de fallar."""
    try:
        import tkinter

        raiz = tkinter.Tk()
    except Exception:  # pragma: no cover - depende de la maquina, no del codigo
        return False
    raiz.destroy()
    return True


class _DialogosCallados:
    """Sustituye los cuadros de dialogo y apunta lo que se les pidio ensenar."""

    def __init__(self):
        self.mensajes = []

    def informar(self, titulo, mensaje, **_):
        self.mensajes.append((titulo, mensaje))

    advertir = informar
    avisar_de_un_error = informar

    def preguntar_si_o_no(self, titulo, mensaje, **_):
        self.mensajes.append((titulo, mensaje))
        return True


@unittest.skipUnless(_hay_ventanas(), "esta maquina no puede abrir ventanas")
class PruebaDeLasPantallasDeLaFase8(unittest.TestCase):
    """Se abre la aplicacion de verdad sobre una carpeta temporal."""

    def setUp(self):
        from interfaz.aplicacion import Aplicacion

        self.carpeta = Path(tempfile.mkdtemp(prefix="fichas_fase8_"))
        self.aplicacion = Aplicacion(carpeta_de_datos=self.carpeta, escribir=lambda _: None)
        self.aplicacion.raiz.withdraw()
        self.conexion = self.aplicacion.conexion

        self.caso = alta_de_caso(
            self.conexion, "CASP2609", unidad_numero="123456",
            fecha_viaje=date.today().isoformat(), estado_recomendacion="incompleta",
        ).id
        self.persona = alta_de_persona(
            self.conexion, self.caso, mrn="055-1111-3853", nombre="José Peña",
            fila_formulario=1,
        )

    def tearDown(self):
        self.conexion.close()
        self.aplicacion.raiz.destroy()
        shutil.rmtree(self.carpeta, ignore_errors=True)

    def _textos(self, widget):
        """Todo el texto dibujado bajo ese widget, mirando hijo por hijo."""
        textos = []
        try:
            textos.append(str(widget.cget("text")))
        except Exception:
            pass
        for hijo in widget.winfo_children():
            textos.extend(self._textos(hijo))
        return textos

    def test_el_inicio_lleva_a_la_pantalla_de_reportes(self):
        """Sin un camino desde el inicio, reportes no existe para quien la usa.

        ⚠️ **Esta prueba cambio el 2026-09-03, y cambio lo que MIRA, no lo que
        exige.** Antes buscaba la palabra «Reportes» dibujada en el inicio, porque
        habia un boton con ese rotulo. El panel adopto los rotulos del mockup v2 y
        ese boton ya no se llama asi: ahora se llega por **«Historial»** en el menu
        de iconos y por **«Informe para la dirección (PDF)»** entre las cuatro
        acciones.

        Buscar un rotulo comprobaba que hubiera una palabra escrita; esto comprueba
        que **el camino lleva de verdad a la pantalla**, que es lo que el titulo de
        la prueba siempre dijo. Un boton bien rotulado que no navega habria pasado
        la version vieja de esta prueba.
        """
        self.aplicacion.mostrar_inicio()
        inicio = self.aplicacion._pantalla

        # El camino del menu de iconos.
        inicio.menu.ir_a("Historial")
        self.assertIsInstance(self.aplicacion._pantalla, PantallaDeReportes)

        # Y el de las cuatro acciones, que tiene que llevar al mismo sitio.
        self.aplicacion.mostrar_inicio()
        self.aplicacion._pantalla._generar_el_informe()
        self.assertIsInstance(self.aplicacion._pantalla, PantallaDeReportes)

    def test_la_pantalla_de_reportes_pinta_las_tres_metricas(self):
        """El criterio de cierre pide que las tres salgan con su numero."""
        self.aplicacion.mostrar_reportes()
        dibujado = " | ".join(self._textos(self.aplicacion.raiz))
        self.assertIn("1. Casos verificados en el período", dibujado)
        self.assertIn("2. Tiempo promedio de importar a verificar", dibujado)
        self.assertIn("3. Casos con el problema detectado ANTES", dibujado)

    def test_la_pantalla_de_reportes_genera_los_dos_archivos(self):
        """El boton escribe el `.xlsx` y el `.pdf` donde dice que los escribe."""
        import interfaz.reportes as pantalla_de_reportes

        dialogos = _DialogosCallados()
        original = pantalla_de_reportes.dialogos
        pantalla_de_reportes.dialogos = dialogos
        try:
            self.aplicacion.mostrar_reportes()
            resultado = self.aplicacion._pantalla.generar()
        finally:
            pantalla_de_reportes.dialogos = original

        self.assertTrue(resultado.excel.ruta.exists())
        self.assertTrue(resultado.pdf.ruta.exists())
        self.assertIn(str(self.carpeta), str(resultado.excel.ruta))

    def test_la_pantalla_de_archivo_archiva_y_anota_el_motivo(self):
        """El camino entero: elegir «no pudo viajar», escribir el motivo, archivar."""
        import interfaz.archivar as pantalla_de_archivo

        dialogos = _DialogosCallados()
        original = pantalla_de_archivo.dialogos
        pantalla_de_archivo.dialogos = dialogos
        try:
            self.aplicacion.archivar_caso(self.caso, "CASP2609")
            bloque = self.aplicacion._pantalla._bloques[0]
            bloque.respuesta.set("no")
            bloque._ajustar_el_motivo()
            bloque.motivo.insert(0, "la recomendación llegó sin firma")
            self.assertTrue(self.aplicacion._pantalla.archivar())
        finally:
            pantalla_de_archivo.dialogos = original

        self.assertTrue(esta_archivado(self.conexion, self.caso))
        fila = self.conexion.execute(
            "SELECT pudo_viajar, motivo_no_viajo FROM personas WHERE id = ?",
            (self.persona,),
        ).fetchone()
        self.assertEqual(fila["pudo_viajar"], NO_PUDO_VIAJAR)
        self.assertIn("sin firma", fila["motivo_no_viajo"])

    def test_archivar_sin_motivo_no_archiva_nada(self):
        """Si el motivo falta, el caso se queda donde estaba. Ni medio archivado."""
        import interfaz.archivar as pantalla_de_archivo

        dialogos = _DialogosCallados()
        original = pantalla_de_archivo.dialogos
        pantalla_de_archivo.dialogos = dialogos
        try:
            self.aplicacion.archivar_caso(self.caso, "CASP2609")
            self.aplicacion._pantalla._bloques[0].respuesta.set("no")
            self.assertFalse(self.aplicacion._pantalla.archivar())
        finally:
            pantalla_de_archivo.dialogos = original

        self.assertFalse(esta_archivado(self.conexion, self.caso))
        self.assertTrue(any("No se pudo archivar" in t for t, _ in dialogos.mensajes))

    def test_el_espejo_se_regenera_al_archivar(self):
        """El Excel espejo tiene columna `archivado`: sin regenerar, mentiria."""
        import interfaz.archivar as pantalla_de_archivo
        from openpyxl import load_workbook

        from espejo.rutas import ruta_del_espejo

        original = pantalla_de_archivo.dialogos
        pantalla_de_archivo.dialogos = _DialogosCallados()
        try:
            self.aplicacion.archivar_caso(self.caso, "CASP2609")
            self.aplicacion._pantalla.archivar()
        finally:
            pantalla_de_archivo.dialogos = original

        hoja = load_workbook(ruta_del_espejo(self.carpeta))["casos"]
        cabecera = [celda.value for celda in next(hoja.iter_rows())]
        valores = dict(zip(cabecera, [c.value for c in list(hoja.iter_rows())[1]]))
        self.assertEqual(valores["archivado"], 1)
        self.assertIsNotNone(valores["fecha_archivado"])

    def test_un_caso_que_no_existe_no_rompe_la_ventana(self):
        """Un id que no esta se dice con un cuadro de error, no con una traza."""
        import interfaz.aplicacion as ventana

        dialogos = _DialogosCallados()
        original = ventana.dialogos
        ventana.dialogos = dialogos
        try:
            self.aplicacion.archivar_caso(9999, "NOEXISTE")
        finally:
            ventana.dialogos = original
        self.assertTrue(dialogos.mensajes)
        self.assertRaises(ErrorDeValidacion, self._archivar_un_caso_que_no_existe)

    def _archivar_un_caso_que_no_existe(self):
        from datos.archivo import archivar_caso

        archivar_caso(self.conexion, 9999)


if __name__ == "__main__":
    unittest.main()
